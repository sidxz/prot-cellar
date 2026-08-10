"""API tests for ImportType.TARGET_BIOLOGY — the workbook import.

Every test drives the same three steps: upload a small in-memory workbook,
POST /api/v1/imports to create the run (exercising the real params validation
and target_key/upload_ref_of routing), then hand the run to the worker
directly (mirroring test_import_worker.py / test_plugin_run.py) since no arq
worker process consumes the queue in tests.
"""

from __future__ import annotations

import io
import random
import uuid
from typing import Any

import openpyxl
import pytest
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from protcellar.domain.imports.enums import ImportStatus
from protcellar.domain.shared.global_workspace import SHARED_WORKSPACE_ID
from protcellar.infrastructure.ingestion import worker as worker_mod
from protcellar.infrastructure.messaging.event_dispatcher import EventDispatcher
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.unpublished_structure_repository import (  # noqa: E501
    SQLAlchemyUnpublishedStructureRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.vulnerability_repository import (  # noqa: E501
    SQLAlchemyVulnerabilityRepository,
)
from protcellar.infrastructure.persistence.unit_of_work import AsyncUnitOfWork

pytestmark = pytest.mark.asyncio


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


def _workbook(sheet_name: str, header: list[str], rows: list[list[Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet_name
    ws.append(header)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _upload(client: AsyncClient, data: bytes) -> str:
    resp = await client.post(
        "/api/v1/imports/uploads",
        params={"import_type": "target_biology"},
        files={
            "file": (
                "workbook.xlsx",
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert resp.status_code == 200, resp.text
    return str(resp.json()["upload_ref"])


async def _start(client: AsyncClient, **params: Any) -> dict[str, Any]:
    resp = await client.post(
        "/api/v1/imports", json={"import_type": "target_biology", "params": params}
    )
    assert resp.status_code == 202, resp.text
    return dict(resp.json())


async def _run_worker(factory: async_sessionmaker, run_id: str, workspace_id: uuid.UUID) -> None:
    ctx = {"session_factory": factory, "dispatcher": EventDispatcher()}
    await worker_mod.run_import(ctx, run_id, str(workspace_id))


async def _organism(client: AsyncClient) -> str:
    tax_id = random.randint(200_000, 999_999)
    resp = await client.post(
        "/api/v1/organisms",
        json={
            "ncbi_tax_id": tax_id,
            "rank": "species",
            "scientific_name": "Mycobacterium tuberculosis",
        },
    )
    assert resp.status_code == 201, resp.text
    return str(resp.json()["id"])


async def _gene(
    client: AsyncClient, organism_id: str, *, locus: str, synonyms: list[str] | None = None
) -> None:
    rec = {
        "primary_name": locus,
        "organism_id": organism_id,
        "source": "test",
        "source_release": "1",
        "source_record_id": f"{organism_id}:{locus}",
        "source_record_checksum": "c1",
    }
    if synonyms:
        rec["synonyms"] = synonyms
    resp = await client.post("/api/v1/genes/bulk", json={"records": [rec]})
    assert resp.status_code == 200, resp.text
    assert resp.json()["summary"]["created"] == 1, resp.json()


async def _protein(client: AsyncClient, organism_id: str, *, accession: str) -> None:
    rec = {
        "primary_accession": accession,
        "organism_id": organism_id,
        "sequence": "MADQLTEEQIAEFKEAFSLF",
        "is_reviewed": True,
        "source": "test",
        "source_release": "1",
        "source_record_id": accession,
        "source_record_checksum": "c1",
    }
    resp = await client.post("/api/v1/proteins/bulk", json={"records": [rec]})
    assert resp.status_code == 200, resp.text
    assert resp.json()["summary"]["created"] == 1, resp.json()


# ---------------------------------------------------------------------------
# tests
# ---------------------------------------------------------------------------


async def test_a_preview_run_commits_nothing(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """Count records before and after. A preview that writes is the one
    unrecoverable bug here."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0001")
        upload_ref = await _upload(
            client, _workbook("vulnerability", ["locus_tag", "condition"], [["Rv0001", "hypoxia"]])
        )

        run = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=True)
        await _run_worker(factory, run["id"], workspace_id)

        got = await client.get(f"/api/v1/imports/{run['id']}")
        body = got.json()
        assert body["status"] == ImportStatus.SUCCEEDED.value
        assert body["summary"]["kinds"]["vulnerability"]["create"] == 1
        # Unlike every other import type, TARGET_BIOLOGY's params must still
        # round-trip — the preview screen's Apply button reads organism_id/
        # match_by/update_existing straight off it.
        assert body["params"]["organism_id"] == organism_id

        async with AsyncUnitOfWork(factory) as uow:
            repo = SQLAlchemyVulnerabilityRepository(uow)
            got_gene = await client.get(f"/api/v1/genes?name=Rv0001&organism_id={organism_id}")
            gene_id = uuid.UUID(got_gene.json()["items"][0]["id"])
            records = await repo.find_by_gene(workspace_id, gene_id)
        assert records == []
    finally:
        await engine.dispose()


async def test_apply_produces_exactly_the_previewed_counts(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0002")
        upload_ref = await _upload(
            client,
            _workbook(
                "vulnerability",
                ["locus_tag", "condition"],
                [["Rv0002", "hypoxia"], ["Rv0002", "aerobic"]],
            ),
        )

        preview = await _start(
            client, upload_ref=upload_ref, organism_id=organism_id, dry_run=True
        )
        await _run_worker(factory, preview["id"], workspace_id)
        preview_summary = (await client.get(f"/api/v1/imports/{preview['id']}")).json()["summary"]
        assert preview_summary["kinds"]["vulnerability"]["create"] == 2

        apply_run = await _start(
            client, upload_ref=upload_ref, organism_id=organism_id, dry_run=False
        )
        assert apply_run["id"] != preview["id"], "apply must be a NEW run, not a mutation"
        await _run_worker(factory, apply_run["id"], workspace_id)
        apply_summary = (await client.get(f"/api/v1/imports/{apply_run['id']}")).json()["summary"]
        assert apply_summary["kinds"]["vulnerability"]["create"] == 2

        async with AsyncUnitOfWork(factory) as uow:
            repo = SQLAlchemyVulnerabilityRepository(uow)
            got_gene = await client.get(f"/api/v1/genes?name=Rv0002&organism_id={organism_id}")
            gene_id = uuid.UUID(got_gene.json()["items"][0]["id"])
            records = await repo.find_by_gene(workspace_id, gene_id)
        assert len(records) == 2
    finally:
        await engine.dispose()


async def test_the_import_lands_in_the_callers_workspace_not_shared(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """This corpus is private. A row in SHARED is readable by every tenant and
    editable by none."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0003")
        upload_ref = await _upload(client, _workbook("vulnerability", ["locus_tag"], [["Rv0003"]]))
        run = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=False)
        await _run_worker(factory, run["id"], workspace_id)

        got_gene = await client.get(f"/api/v1/genes?name=Rv0003&organism_id={organism_id}")
        gene_id = got_gene.json()["items"][0]["id"]

        bundle = await client.get(f"/api/v1/genes/{gene_id}/target-biology")
        assert len(bundle.json()["vulnerability"]) == 1

        async with AsyncUnitOfWork(factory) as uow:
            repo = SQLAlchemyVulnerabilityRepository(uow)
            records = await repo.find_by_gene(workspace_id, uuid.UUID(gene_id))
        assert len(records) == 1
        assert records[0].workspace_id == workspace_id
        assert records[0].workspace_id != SHARED_WORKSPACE_ID
    finally:
        await engine.dispose()


async def test_a_client_supplied_target_workspace_id_is_ignored(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """TargetBiologyParams has no target_workspace_id field at all — the value
    always comes from the run's own workspace_id (StartImport's auth.workspace_id
    at creation time), never the request body. A stray field here must be inert,
    not a cross-tenant write."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        smuggled_workspace_id = uuid.uuid4()
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0011")
        upload_ref = await _upload(client, _workbook("vulnerability", ["locus_tag"], [["Rv0011"]]))
        run = await _start(
            client,
            upload_ref=upload_ref,
            organism_id=organism_id,
            dry_run=False,
            target_workspace_id=str(smuggled_workspace_id),
        )
        await _run_worker(factory, run["id"], workspace_id)

        got_gene = await client.get(f"/api/v1/genes?name=Rv0011&organism_id={organism_id}")
        gene_id = uuid.UUID(got_gene.json()["items"][0]["id"])

        async with AsyncUnitOfWork(factory) as uow:
            repo = SQLAlchemyVulnerabilityRepository(uow)
            mine = await repo.find_by_gene(workspace_id, gene_id)
            smuggled = await repo.find_by_gene(smuggled_workspace_id, gene_id)
        assert len(mine) == 1
        assert mine[0].workspace_id == workspace_id
        assert smuggled == []
    finally:
        await engine.dispose()


async def test_a_second_workspace_sees_none_of_it(
    client: AsyncClient,
    other_workspace_client: AsyncClient,
    database_url: str,
    _run_migrations: None,
    workspace_id: uuid.UUID,
) -> None:
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0004")
        upload_ref = await _upload(client, _workbook("vulnerability", ["locus_tag"], [["Rv0004"]]))
        run = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=False)
        await _run_worker(factory, run["id"], workspace_id)

        got_gene = await client.get(f"/api/v1/genes?name=Rv0004&organism_id={organism_id}")
        gene_id = got_gene.json()["items"][0]["id"]

        mine = await client.get(f"/api/v1/genes/{gene_id}/target-biology")
        assert len(mine.json()["vulnerability"]) == 1

        theirs = await other_workspace_client.get(f"/api/v1/genes/{gene_id}/target-biology")
        assert theirs.json()["vulnerability"] == []
    finally:
        await engine.dispose()


async def test_an_unmatched_locus_fails_its_row_and_the_rest_still_import(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0005")
        upload_ref = await _upload(
            client,
            _workbook("vulnerability", ["locus_tag"], [["Rv0005"], ["Rv9999-does-not-exist"]]),
        )
        run = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=True)
        await _run_worker(factory, run["id"], workspace_id)

        summary = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]
        kind = summary["kinds"]["vulnerability"]
        assert kind["create"] == 1
        assert kind["failed"] == 1
        assert summary["unmatched"]["count"] == 1
        assert "Rv9999-does-not-exist" in summary["unmatched"]["examples"]
        assert any("Rv9999-does-not-exist" in p["reason"] for p in summary["problems"])
    finally:
        await engine.dispose()


async def test_an_ambiguous_gene_name_fails_its_row_and_names_the_candidates(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """build_locus_index used to use setdefault, so the first match silently
    won. Tolerable for a code-driven CLI, not for an operator's spreadsheet."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0006", synonyms=["sharedName"])
        await _gene(client, organism_id, locus="Rv0007", synonyms=["sharedName"])
        upload_ref = await _upload(
            client, _workbook("vulnerability", ["gene_name"], [["sharedName"]])
        )
        run = await _start(
            client,
            upload_ref=upload_ref,
            organism_id=organism_id,
            match_by="gene_name",
            dry_run=True,
        )
        await _run_worker(factory, run["id"], workspace_id)

        summary = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]
        kind = summary["kinds"]["vulnerability"]
        assert kind["create"] == 0
        assert kind["failed"] == 1
        reasons = [p["reason"] for p in summary["problems"]]
        matching = next(r for r in reasons if "ambiguous" in r.lower())
        assert "Rv0006" in matching
        assert "Rv0007" in matching
    finally:
        await engine.dispose()


async def test_a_real_row_number_reaches_the_response_but_a_workbook_level_one_stays_none(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """`parse_workbook` puts a real spreadsheet row on a problem raised while
    walking actual rows (an unrecognised column here), and a placeholder
    ``row=1`` on a workbook-level one (no locus_tag column at all) — the
    latter is not a row an operator should go check, so it must not reach the
    API pretending to be one."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "vulnerability"
        ws.append(["locus_tag", "bogus_column"])
        ws.append(["Rv0099", "x"])
        hypomorph = wb.create_sheet("hypomorph")
        hypomorph.append(["condition"])  # no locus_tag column at all
        hypomorph.append(["stress"])
        buf = io.BytesIO()
        wb.save(buf)
        upload_ref = await _upload(client, buf.getvalue())

        run = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=True)
        await _run_worker(factory, run["id"], workspace_id)

        problems = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]["problems"]

        real_row = next(p for p in problems if "unrecognised column" in p["reason"])
        assert real_row["row"] == 2

        workbook_level = next(p for p in problems if "no 'locus_tag' column found" in p["reason"])
        assert workbook_level["row"] is None
    finally:
        await engine.dispose()


async def test_ignored_columns_names_the_dropped_provenance_headers(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0008")
        upload_ref = await _upload(
            client,
            _workbook(
                "vulnerability",
                ["locus_tag", "contributor", "note"],
                [["Rv0008", "J. Doe", "from a screen"]],
            ),
        )
        run = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=True)
        await _run_worker(factory, run["id"], workspace_id)

        summary = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]
        assert summary["kinds"]["vulnerability"]["create"] == 1
        assert sorted(summary["ignored_columns"]["vulnerability"]) == ["contributor", "note"]
    finally:
        await engine.dispose()


async def test_already_present_counts_before_this_runs_own_writes(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0009")
        upload_ref = await _upload(client, _workbook("vulnerability", ["locus_tag"], [["Rv0009"]]))

        first = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=False)
        await _run_worker(factory, first["id"], workspace_id)
        first_summary = (await client.get(f"/api/v1/imports/{first['id']}")).json()["summary"]
        assert first_summary["already_present"]["vulnerability"] == 0

        second = await _start(
            client, upload_ref=upload_ref, organism_id=organism_id, dry_run=False
        )
        await _run_worker(factory, second["id"], workspace_id)
        second_summary = (await client.get(f"/api/v1/imports/{second['id']}")).json()["summary"]
        # Add mode (the default): the bulk command never sees the first run's
        # row as "existing" (see _NoExistingMatch), so the second run creates a
        # second one — already_present warns about exactly this before it runs.
        assert second_summary["already_present"]["vulnerability"] == 1
        assert second_summary["kinds"]["vulnerability"]["create"] == 1

        got_gene = await client.get(f"/api/v1/genes?name=Rv0009&organism_id={organism_id}")
        gene_id = got_gene.json()["items"][0]["id"]
        async with AsyncUnitOfWork(factory) as uow:
            repo = SQLAlchemyVulnerabilityRepository(uow)
            records = await repo.find_by_gene(workspace_id, uuid.UUID(gene_id))
        assert len(records) == 2, "add mode run twice doubles the data"
    finally:
        await engine.dispose()


async def test_update_mode_matches_instead_of_duplicating(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """The mirror of the add-mode doubling case: update_existing=True lets the
    bulk command's own natural-key match run, so re-importing the same row
    updates in place rather than creating a second one."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0010")
        upload_ref = await _upload(client, _workbook("vulnerability", ["locus_tag"], [["Rv0010"]]))

        first = await _start(
            client,
            upload_ref=upload_ref,
            organism_id=organism_id,
            update_existing=True,
            dry_run=False,
        )
        await _run_worker(factory, first["id"], workspace_id)

        second = await _start(
            client,
            upload_ref=upload_ref,
            organism_id=organism_id,
            update_existing=True,
            dry_run=False,
        )
        await _run_worker(factory, second["id"], workspace_id)
        second_summary = (await client.get(f"/api/v1/imports/{second['id']}")).json()["summary"]
        assert second_summary["kinds"]["vulnerability"]["update"] == 1
        assert second_summary["kinds"]["vulnerability"]["create"] == 0

        got_gene = await client.get(f"/api/v1/genes?name=Rv0010&organism_id={organism_id}")
        gene_id = got_gene.json()["items"][0]["id"]
        async with AsyncUnitOfWork(factory) as uow:
            repo = SQLAlchemyVulnerabilityRepository(uow)
            records = await repo.find_by_gene(workspace_id, uuid.UUID(gene_id))
        assert len(records) == 1
    finally:
        await engine.dispose()


async def test_already_present_ignores_unrelated_shared_reference_rows(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """already_present must count only what the target workspace itself owns.
    list_paginated is readable_by (target OR SHARED) — a brand-new tenant with
    zero rows of its own must not see a nonzero count just because SHARED
    holds an unrelated legacy row for the same gene."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0012")
        got_gene = await client.get(f"/api/v1/genes?name=Rv0012&organism_id={organism_id}")
        gene_id = uuid.UUID(got_gene.json()["items"][0]["id"])

        # A SHARED vulnerability row for the same gene — legacy reference data,
        # not this tenant's, and not something add mode run twice could ever
        # double (the workbook import never writes SHARED).
        async with AsyncUnitOfWork(factory) as uow:
            from protcellar.domain.shared.provenance import Provenance, ProvenanceSourceType
            from protcellar.domain.target_biology.vulnerability import Vulnerability

            shared_record = Vulnerability.create(
                workspace_id=SHARED_WORKSPACE_ID,
                gene_id=gene_id,
                provenance=Provenance(source_type=ProvenanceSourceType.PUBLISHED),
                vulnerability_score=0.5,
            )
            await SQLAlchemyVulnerabilityRepository(uow).save(shared_record)
            await uow.commit()

        upload_ref = await _upload(client, _workbook("vulnerability", ["locus_tag"], [["Rv0012"]]))
        run = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=True)
        await _run_worker(factory, run["id"], workspace_id)

        summary = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]
        assert summary["already_present"]["vulnerability"] == 0
    finally:
        await engine.dispose()


async def test_update_mode_warns_about_unresolved_ligand_text(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """unpublished_structure's update-mode key is (protein_id, method,
    ligands), and ligands only ever means resolved compound ids. Two rows for
    the same protein/method whose ligand text never resolved are
    indistinguishable to that key — the second silently overwrites the first.
    The preview must say so.

    The warning is computed straight off the parsed rows (extensions
    carrying ligand_reported), so it fires under dry_run too — but the count
    collapse it describes is an *apply*-only observation: under dry_run the
    bulk command never saves row 1, so row 2's own "does a match already
    exist" lookup (a DB query) can't see it either, and both come back
    "created". Only a real apply serializes the two rows against the database
    in between, which is when the second one finds and overwrites the first.
    """
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _protein(client, organism_id, accession="P0DP99")
        upload_ref = await _upload(
            client,
            _workbook(
                "unpublished_structure",
                ["accession", "method", "ligand_ids"],
                [
                    ["P0DP99", "X-ray", "Apo"],
                    ["P0DP99", "X-ray", "SO4 bound"],
                ],
            ),
        )
        run = await _start(
            client,
            upload_ref=upload_ref,
            organism_id=organism_id,
            update_existing=True,
            dry_run=False,
        )
        await _run_worker(factory, run["id"], workspace_id)

        summary = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]
        warning = next(w for w in summary["warnings"] if w["sheet"] == "unpublished_structure")
        assert warning["count"] == 2

        # The bug itself: both rows share the command's match key (method +
        # empty resolved-ligands), so the second updates the first in place
        # instead of creating a second, distinct structure.
        kind = summary["kinds"]["unpublished_structure"]
        assert kind["create"] == 1
        assert kind["update"] == 1
        assert kind["failed"] == 0

        async with AsyncUnitOfWork(factory) as uow:
            got_protein = await client.get("/api/v1/proteins/P0DP99")
            protein_id = uuid.UUID(got_protein.json()["id"])
            records = await SQLAlchemyUnpublishedStructureRepository(uow).find_by_protein(
                workspace_id, protein_id
            )
        assert len(records) == 1, "the two distinct structures collapsed into one"
    finally:
        await engine.dispose()


async def test_update_mode_warns_about_a_natural_key_collision_within_one_sheet(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """Two rows in one sheet sharing the kind's upsert key (condition + method,
    here) are indistinguishable to preview — dry_run never saves, so row 2's
    own "does a match exist" lookup can't see row 1 — but apply's row 1 save()
    autoflushes, so row 2 matches and updates it in place instead of creating
    a second record. The preview must warn, even though (without touching the
    eight commands) its create/update split still won't equal apply's."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0013")
        upload_ref = await _upload(
            client,
            _workbook(
                "vulnerability",
                ["locus_tag", "condition", "method", "vulnerability_score"],
                [
                    ["Rv0013", "hypoxia", "CRISPRi", "0.5"],
                    ["Rv0013", "hypoxia", "CRISPRi", "0.9"],
                ],
            ),
        )

        preview = await _start(
            client,
            upload_ref=upload_ref,
            organism_id=organism_id,
            update_existing=True,
            dry_run=True,
        )
        await _run_worker(factory, preview["id"], workspace_id)
        preview_summary = (await client.get(f"/api/v1/imports/{preview['id']}")).json()["summary"]
        warning = next(w for w in preview_summary["warnings"] if w["sheet"] == "vulnerability")
        assert warning["count"] == 2
        assert preview_summary["kinds"]["vulnerability"]["create"] == 2
        assert preview_summary["kinds"]["vulnerability"]["update"] == 0

        apply_run = await _start(
            client,
            upload_ref=upload_ref,
            organism_id=organism_id,
            update_existing=True,
            dry_run=False,
        )
        await _run_worker(factory, apply_run["id"], workspace_id)
        apply_summary = (await client.get(f"/api/v1/imports/{apply_run['id']}")).json()["summary"]
        assert apply_summary["kinds"]["vulnerability"]["create"] == 1
        assert apply_summary["kinds"]["vulnerability"]["update"] == 1

        got_gene = await client.get(f"/api/v1/genes?name=Rv0013&organism_id={organism_id}")
        gene_id = uuid.UUID(got_gene.json()["items"][0]["id"])
        async with AsyncUnitOfWork(factory) as uow:
            repo = SQLAlchemyVulnerabilityRepository(uow)
            records = await repo.find_by_gene(workspace_id, gene_id)
        assert len(records) == 1, "the second row overwrote the first in place"
    finally:
        await engine.dispose()


async def test_update_mode_collision_warning_never_fires_in_add_mode(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """Add mode never matches at all (see _NoExistingMatch), so two rows
    sharing a natural key just become two separate records — nothing for the
    warning to say."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0015")
        upload_ref = await _upload(
            client,
            _workbook(
                "vulnerability",
                ["locus_tag", "condition", "method", "vulnerability_score"],
                [
                    ["Rv0015", "hypoxia", "CRISPRi", "0.5"],
                    ["Rv0015", "hypoxia", "CRISPRi", "0.9"],
                ],
            ),
        )
        run = await _start(
            client, upload_ref=upload_ref, organism_id=organism_id, dry_run=True
        )  # update_existing defaults to False
        await _run_worker(factory, run["id"], workspace_id)

        summary = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]
        assert summary["warnings"] == []
        assert summary["kinds"]["vulnerability"]["create"] == 2
    finally:
        await engine.dispose()


async def test_hypomorph_resolves_a_same_workbook_strain_regardless_of_sheet_order(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """crispri_strain must dispatch before hypomorph even when the workbook
    lists hypomorph's sheet first — otherwise apply's result depends on an
    operator's sheet order, which a dry-run preview (which persists nothing,
    so it always fails to resolve a same-workbook strain either way) could
    never foretell."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0014")

        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "hypomorph"  # listed BEFORE crispri_strain in the workbook
        ws.append(["locus_tag", "growth_defect", "knockdown_strain"])
        ws.append(["Rv0014", "true", "strainA"])
        strains = wb.create_sheet("crispri_strain")
        strains.append(["locus_tag", "name"])
        strains.append(["Rv0014", "strainA"])
        buf = io.BytesIO()
        wb.save(buf)
        upload_ref = await _upload(client, buf.getvalue())

        run = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=False)
        await _run_worker(factory, run["id"], workspace_id)

        summary = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]
        assert summary["kinds"]["crispri_strain"]["create"] == 1
        assert summary["kinds"]["hypomorph"]["create"] == 1
        assert summary["kinds"]["hypomorph"]["failed"] == 0
    finally:
        await engine.dispose()


async def test_two_sheets_normalising_to_the_same_kind_only_the_first_is_used(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """ "vulnerability" and "Vulnerability " both normalise to the same kind.
    Before the fix, one plan existed per *sheet*, so both got dispatched and
    written while kinds_summary/already_present (assigned, not accumulated,
    keyed by kind) silently kept only the last one's numbers. Now the second
    sheet is a workbook-level problem and is never parsed at all — only the
    first sheet's row lands."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _gene(client, organism_id, locus="Rv0016")
        await _gene(client, organism_id, locus="Rv0017")

        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "vulnerability"
        ws.append(["locus_tag", "condition"])
        ws.append(["Rv0016", "hypoxia"])
        dup = wb.create_sheet("Vulnerability ")
        dup.append(["locus_tag", "condition"])
        dup.append(["Rv0017", "normoxia"])
        buf = io.BytesIO()
        wb.save(buf)
        upload_ref = await _upload(client, buf.getvalue())

        run = await _start(client, upload_ref=upload_ref, organism_id=organism_id, dry_run=True)
        await _run_worker(factory, run["id"], workspace_id)

        summary = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]
        assert list(summary["kinds"]) == ["vulnerability"]
        assert summary["kinds"]["vulnerability"]["create"] == 1
        assert summary["already_present"] == {"vulnerability": 0}
        assert any("duplicate sheet" in p["reason"] for p in summary["problems"])
    finally:
        await engine.dispose()


async def test_add_mode_never_triggers_the_ligand_warning(
    client: AsyncClient, database_url: str, _run_migrations: None, workspace_id: uuid.UUID
) -> None:
    """Add mode never matches at all (see _NoExistingMatch) — the collapse the
    warning describes can't happen there, so it must not fire."""
    engine = create_async_engine(database_url)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    try:
        organism_id = await _organism(client)
        await _protein(client, organism_id, accession="P0DP98")
        upload_ref = await _upload(
            client,
            _workbook(
                "unpublished_structure",
                ["accession", "method", "ligand_ids"],
                [["P0DP98", "X-ray", "Apo"], ["P0DP98", "X-ray", "SO4 bound"]],
            ),
        )
        run = await _start(
            client, upload_ref=upload_ref, organism_id=organism_id, dry_run=True
        )  # update_existing defaults to False
        await _run_worker(factory, run["id"], workspace_id)

        summary = (await client.get(f"/api/v1/imports/{run['id']}")).json()["summary"]
        assert summary["warnings"] == []
        assert summary["kinds"]["unpublished_structure"]["create"] == 2
    finally:
        await engine.dispose()
