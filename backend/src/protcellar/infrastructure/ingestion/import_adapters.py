"""Cross-context runner wiring — the infrastructure adapter layer for imports.

Each :class:`ImportAdapter` builds a fresh :class:`AsyncUnitOfWork`, wires
repos / bulk use-cases / HTTP clients, and delegates to the matching runner —
exactly as the CLI scripts do, but using the :class:`ImportRuntime` injected by
the background worker instead of local engine + noop dispatcher.

``IMPORT_ADAPTERS`` is the registry the worker calls by :class:`ImportType`.

**IMPORTANT**: All three runner classes are imported at module top level so that
monkeypatching (e.g. ``protcellar.infrastructure.ingestion.import_adapters.ProteomeImportRunner``)
resolves correctly in tests.
"""

from __future__ import annotations

import dataclasses
import io
import os
import uuid
from collections.abc import Awaitable, Callable, Sequence
from dataclasses import dataclass
from typing import Any, Protocol

import httpx
import openpyxl
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from protcellar.application.auth import AuthContext
from protcellar.application.imports.params import TargetBiologyParams
from protcellar.application.imports.progress_reporter import ProgressReporter
from protcellar.application.plugins.context import PluginRunContext
from protcellar.application.protein_catalog.bulk_enrich_genes import BulkEnrichGenes
from protcellar.application.protein_catalog.bulk_upsert_genes import BulkUpsertGenes
from protcellar.application.protein_catalog.bulk_upsert_proteins import BulkUpsertProteins
from protcellar.application.shared.event_dispatcher import EventDispatcherProtocol
from protcellar.application.target_biology._import_support import (
    ItemResult,
    LocusIndex,
    build_locus_index,
)
from protcellar.application.target_biology.bulk_upsert_crispri_strain import (
    BulkUpsertCrispriStrain,
    BulkUpsertCrispriStrainCommand,
)
from protcellar.application.target_biology.bulk_upsert_essentiality import (
    BulkUpsertEssentiality,
    BulkUpsertEssentialityCommand,
)
from protcellar.application.target_biology.bulk_upsert_hypomorph import (
    BulkUpsertHypomorph,
    BulkUpsertHypomorphCommand,
)
from protcellar.application.target_biology.bulk_upsert_protein_activity_assay import (
    BulkUpsertProteinActivityAssay,
    BulkUpsertProteinActivityAssayCommand,
)
from protcellar.application.target_biology.bulk_upsert_protein_production import (
    BulkUpsertProteinProduction,
    BulkUpsertProteinProductionCommand,
)
from protcellar.application.target_biology.bulk_upsert_resistance_mutation import (
    BulkUpsertResistanceMutation,
    BulkUpsertResistanceMutationCommand,
)
from protcellar.application.target_biology.bulk_upsert_unpublished_structure import (
    BulkUpsertUnpublishedStructure,
    BulkUpsertUnpublishedStructureCommand,
)
from protcellar.application.target_biology.bulk_upsert_vulnerability import (
    BulkUpsertVulnerability,
    BulkUpsertVulnerabilityCommand,
)
from protcellar.application.target_biology.crud import RecordKind
from protcellar.domain.imports.enums import ImportType
from protcellar.domain.shared.global_workspace import SHARED_WORKSPACE_ID
from protcellar.infrastructure.ingestion.gene_enrichment_runner import GeneEnrichmentRunner
from protcellar.infrastructure.ingestion.go_import_runner import GoImportRunner
from protcellar.infrastructure.ingestion.import_runner import ProteomeImportRunner
from protcellar.infrastructure.ingestion.mycobrowser_client import (
    MYCOBROWSER_H37RV_GFF_URL,
    MycobrowserClient,
)
from protcellar.infrastructure.ingestion.organism_resolver import resolve_organism_id
from protcellar.infrastructure.ingestion.target_biology_workbook import (
    _PROVENANCE_DROPPED,  # reused, not duplicated — see _ignored_columns
    parse_workbook,
)
from protcellar.infrastructure.ingestion.uniprot_client import UniProtClient
from protcellar.infrastructure.ingestion.url_guard import validate_public_url
from protcellar.infrastructure.persistence.sqlalchemy.protein_catalog.gene_repository import (
    SQLAlchemyGeneRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.protein_catalog.protein_repository import (
    SQLAlchemyProteinRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.crispri_strain_repository import (  # noqa: E501
    SQLAlchemyCrispriStrainRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.essentiality_repository import (  # noqa: E501
    SQLAlchemyEssentialityRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.hypomorph_repository import (
    SQLAlchemyHypomorphRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.protein_activity_assay_repository import (  # noqa: E501
    SQLAlchemyProteinActivityAssayRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.protein_production_repository import (  # noqa: E501
    SQLAlchemyProteinProductionRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.resistance_mutation_repository import (  # noqa: E501
    SQLAlchemyResistanceMutationRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.unpublished_structure_repository import (  # noqa: E501
    SQLAlchemyUnpublishedStructureRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.target_biology.vulnerability_repository import (  # noqa: E501
    SQLAlchemyVulnerabilityRepository,
)
from protcellar.infrastructure.persistence.sqlalchemy.workspace_config.extension_field_def_repository import (  # noqa: E501
    SQLAlchemyExtensionFieldDefRepository,
)
from protcellar.infrastructure.persistence.unit_of_work import AsyncUnitOfWork
from protcellar.infrastructure.plugins.in_tree_sink import InTreeSink
from protcellar.infrastructure.plugins.registry import get_plugin

_UNIPROT_BASE_URL = "https://rest.uniprot.org"


# ---------------------------------------------------------------------------
# ImportRuntime — injected by the background worker
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ImportRuntime:
    """All cross-cutting dependencies the worker hands to an adapter."""

    session_factory: async_sessionmaker[AsyncSession]
    dispatcher: EventDispatcherProtocol
    reporter: ProgressReporter
    params: dict[str, Any]
    auth: AuthContext
    load_upload: Callable[[uuid.UUID], Awaitable[bytes]]
    # The run's OWN workspace — set at StartImport time from the caller's real
    # auth.workspace_id, threaded through the arq job payload exactly like the
    # reporter's workspace_id already is (worker.py has no per-request auth of
    # its own; `auth` above is ServiceAuth, which is always SHARED_WORKSPACE_ID
    # and satisfies only the bulk commands' internal require_admin check — it is
    # NOT a tenant). TargetBiologyAdapter is the first adapter that needs to know
    # which tenant a run belongs to, so this is what it reads target_workspace_id
    # from — never rt.auth, never params. Required, not defaulted: a silently
    # SHARED-defaulting tenancy field is exactly the shape of bug Task 1 spent
    # its whole scope removing from the eight bulk commands. worker.py is the
    # only real caller and already passes it; the one test that constructs
    # ImportRuntime directly (test_import_adapters.py) now must too.
    workspace_id: uuid.UUID
    # ponytail: defaulted so legacy adapters/tests need no change; the worker
    # always injects the real ImportRun id for plugin lineage.
    run_id: uuid.UUID = dataclasses.field(default_factory=uuid.uuid4)


# ---------------------------------------------------------------------------
# ImportAdapter protocol
# ---------------------------------------------------------------------------


class ImportAdapter(Protocol):
    """A runnable import strategy keyed on an :class:`ImportType`."""

    import_type: ImportType

    async def run(self, rt: ImportRuntime) -> dict[str, Any]: ...


# ---------------------------------------------------------------------------
# Concrete adapters
# ---------------------------------------------------------------------------


class ProteomeAdapter:
    """Wires :class:`ProteomeImportRunner` from ``rt``."""

    import_type = ImportType.PROTEOME

    async def run(self, rt: ImportRuntime) -> dict[str, Any]:
        params = rt.params
        proteome_id: str = params["proteome_id"]
        force: bool = bool(params.get("force", False))
        dry_run: bool = bool(params.get("dry_run", False))
        limit: int | None = params.get("limit")

        uow = AsyncUnitOfWork(rt.session_factory)
        protein_repo = SQLAlchemyProteinRepository(uow)
        gene_repo = SQLAlchemyGeneRepository(uow)
        bulk = BulkUpsertProteins(uow, protein_repo, rt.dispatcher)
        gene_bulk = BulkUpsertGenes(uow, gene_repo, rt.dispatcher)

        async with httpx.AsyncClient(base_url=_UNIPROT_BASE_URL, timeout=120.0) as http:
            runner = ProteomeImportRunner(
                uow,
                UniProtClient(http),
                bulk,
                gene_bulk=gene_bulk,
                reporter=rt.reporter,
            )
            summary = await runner.run(
                proteome_id,
                dry_run=dry_run,
                limit=limit,
                force=force,
                auth=rt.auth,
            )
        return dataclasses.asdict(summary)


class GeneEnrichmentAdapter:
    """Wires :class:`GeneEnrichmentRunner` from ``rt``."""

    import_type = ImportType.GENE_ENRICHMENT

    async def run(self, rt: ImportRuntime) -> dict[str, Any]:
        params = rt.params

        # Coerce organism_id from JSON string to UUID if provided
        raw_organism_id = params.get("organism_id")
        organism_id: uuid.UUID | None = (
            uuid.UUID(raw_organism_id) if raw_organism_id is not None else None
        )
        tax_id: int = int(params["tax_id"]) if params.get("tax_id") is not None else 83332

        uow = AsyncUnitOfWork(rt.session_factory)
        resolved_id, gene_count = await resolve_organism_id(
            uow, organism_id=organism_id, tax_id=tax_id
        )
        if gene_count == 0:
            raise ValueError(f"organism {resolved_id} has no genes to enrich")

        gff_url: str = params.get("gff_url") or MYCOBROWSER_H37RV_GFF_URL
        # Guard admin-supplied GFF URL against SSRF before any network I/O
        validate_public_url(gff_url)

        uow2 = AsyncUnitOfWork(rt.session_factory)
        gene_repo = SQLAlchemyGeneRepository(uow2)
        bulk = BulkEnrichGenes(uow2, gene_repo, rt.dispatcher)

        async with httpx.AsyncClient() as http:
            client = MycobrowserClient(http)
            runner = GeneEnrichmentRunner(
                bulk,
                client,
                gff_url=gff_url,
                # Use the SSRF-safe fetch so every redirect hop is validated
                gff_fetch=client.fetch_text_secure,
                reporter=rt.reporter,
            )
            summary = await runner.run(resolved_id, auth=rt.auth)
        return dataclasses.asdict(summary)


class GoOntologyAdapter:
    """Wires :class:`GoImportRunner` from ``rt``."""

    import_type = ImportType.GO_ONTOLOGY

    async def run(self, rt: ImportRuntime) -> dict[str, Any]:
        force: bool = bool(rt.params.get("force", False))
        uow = AsyncUnitOfWork(rt.session_factory)
        runner = GoImportRunner(uow, reporter=rt.reporter)
        summary = await runner.run(force=force)
        return dataclasses.asdict(summary)


def _summarize(results: list[ItemResult]) -> dict[str, Any]:
    summary: dict[str, Any] = {"created": 0, "updated": 0, "skipped": 0, "failed": 0}
    for r in results:
        if r.status in summary:
            summary[r.status] += 1
    summary["total"] = len(results)
    return summary


class PluginDispatchAdapter:
    """ImportType.PLUGIN -> resolve rt.params['plugin_id'] from the registry, wire the
    in-tree sink, run the plugin, and summarize. The worker's IMPORT_ADAPTERS lookup
    dispatches here with no plugin-specific branching in worker.py."""

    import_type = ImportType.PLUGIN

    async def run(self, rt: ImportRuntime) -> dict[str, Any]:
        params = rt.params
        plugin = get_plugin(str(params["plugin_id"]))
        manifest = plugin.manifest()
        organism_id = uuid.UUID(str(params["organism_id"])) if params.get("organism_id") else None
        dry_run = bool(params.get("dry_run", False))
        generation_method = manifest.default_generation_method.value

        uow = AsyncUnitOfWork(rt.session_factory)
        gene_repo = SQLAlchemyGeneRepository(uow)
        ess_repo = SQLAlchemyEssentialityRepository(uow)
        essentiality = BulkUpsertEssentiality(uow, gene_repo, ess_repo, rt.dispatcher)

        async def _upsert_essentiality(records: Sequence[object]) -> list[ItemResult]:
            if organism_id is None:
                raise ValueError("essentiality upsert requires an organism_id")
            cmd = BulkUpsertEssentialityCommand(
                target_workspace_id=SHARED_WORKSPACE_ID,
                organism_id=organism_id,
                records=tuple(records),  # type: ignore[arg-type]  # elements are EssentialityImportRecord
                generation_method=generation_method,
                source_run_id=rt.run_id,
                dry_run=dry_run,
            )
            return (await essentiality(cmd, rt.auth)).unwrap()

        # record_type -> upserter. Add an entry to support a new target record; the
        # 7 other BulkUpsert<X> commands already exist in application/target_biology.
        sink = InTreeSink({"essentiality": _upsert_essentiality})

        # Secrets come from the process env by name, never from persisted params.
        secrets = {k: os.environ[k] for k in manifest.requires_secrets if k in os.environ}

        ctx = PluginRunContext(
            params=params,
            organism_id=organism_id,
            load_upload=rt.load_upload,
            sink=sink,
            reporter=rt.reporter,
            auth=rt.auth,
            secrets=secrets,
        )
        await plugin.run(ctx)
        return _summarize(sink.results)


# ---------------------------------------------------------------------------
# TargetBiologyAdapter — ImportType.TARGET_BIOLOGY
# ---------------------------------------------------------------------------

_GENE_SIDE_KINDS: frozenset[RecordKind] = frozenset(
    {
        RecordKind.ESSENTIALITY,
        RecordKind.VULNERABILITY,
        RecordKind.HYPOMORPH,
        RecordKind.CRISPRI_STRAIN,
        RecordKind.RESISTANCE_MUTATION,
    }
)


class _NoExistingMatch:
    """Wraps a target-biology repo so its own-kind "find an existing row to
    match" lookup always reports none — the seam that gives ``update_existing=
    False`` (add mode, this adapter's default) real bypass-the-match semantics
    without touching any of the eight bulk commands, none of which takes a
    parameter to suppress their internal match-then-update. Every command calls
    exactly one of the two methods below, exactly once per row, purely to find
    something to merge into; forcing an empty result routes every row through
    the command's own create branch instead — see the "Record identity, add
    mode" row of the design's Decisions table.

    Everything else — ``save()``, and a *different* repo's own
    ``find_by_gene``/``find_by_protein`` used to resolve a foreign reference
    (hypomorph resolving ``knockdown_strain`` by name against the CrispriStrain
    repo is the one case of this in the eight commands) — passes straight
    through untouched; only the two "owned" match lookups are ever intercepted.
    """

    def __init__(self, inner: Any) -> None:
        self._inner = inner

    async def find_owned_by_gene(self, *_args: Any, **_kwargs: Any) -> list[Any]:
        return []

    async def find_owned_by_protein(self, *_args: Any, **_kwargs: Any) -> list[Any]:
        return []

    def __getattr__(self, name: str) -> Any:
        return getattr(self._inner, name)


@dataclass(frozen=True, kw_only=True)
class _KindContext:
    """Shared per-run context every ``_run_<kind>`` dispatcher closes over."""

    uow: AsyncUnitOfWork
    dispatcher: EventDispatcherProtocol
    auth: AuthContext
    target_workspace_id: uuid.UUID
    organism_id: uuid.UUID
    dry_run: bool
    update_existing: bool
    gene_repo: Any
    protein_repo: Any


async def _already_present(repo: Any, ctx: _KindContext) -> int:
    """How many records of this kind the target workspace itself already owns,
    for the target organism — so the preview can warn that add mode run twice
    doubles *this tenant's own* data (add mode never checks for a match; see
    ``_NoExistingMatch`` above).

    ``list_paginated`` is ``readable_by`` (target workspace OR SHARED), so its
    raw result includes shared reference rows this import will never touch and
    add mode could never duplicate — a brand-new tenant with zero rows of its
    own would otherwise see a nonzero count from unrelated legacy SHARED data
    and learn to ignore the warning. Filtered here to rows the target
    workspace itself owns, rather than adding an ``owned_by`` variant to any
    of the eight repos (outside this task's file list).

    ponytail: counts by fetching every readable row and taking ``len()`` — none
    of the eight repos exposes a dedicated COUNT query (adding one is outside
    this task's file list). Fine at this corpus's scale (tens of thousands of
    rows, one fetch per kind per run, mirroring the cost ``build_locus_index``'s
    own callers already each pay); upgrade to a real ``COUNT(*)`` if a kind's
    table grows past what fits comfortably in memory.

    Runs its own ``async with ctx.uow:`` — the repository is a normal one, so a
    call outside an active UnitOfWork raises ``RuntimeError`` at runtime while
    mypy stays silent (this exact mistake has already been made once in this
    codebase; see ``_require_readable_protein`` in ``interface/routes/target_biology.py``
    for the correct shape).
    """
    async with ctx.uow:
        rows = await repo.list_paginated(
            ctx.target_workspace_id, organism_id=ctx.organism_id, limit=None
        )
    return sum(1 for r in rows if r.workspace_id == ctx.target_workspace_id)


def _match_repo(repo: Any, ctx: _KindContext) -> Any:
    """The repo a bulk command should see for its own "existing row" lookup:
    the real one in update mode, a wrapper that reports none in add mode."""
    return repo if ctx.update_existing else _NoExistingMatch(repo)


async def _run_essentiality(ctx: _KindContext, records: list[Any]) -> tuple[list[ItemResult], int]:
    repo = SQLAlchemyEssentialityRepository(ctx.uow)
    present = await _already_present(repo, ctx)
    cmd = BulkUpsertEssentialityCommand(
        target_workspace_id=ctx.target_workspace_id,
        organism_id=ctx.organism_id,
        records=tuple(records),
        dry_run=ctx.dry_run,
    )
    handler = BulkUpsertEssentiality(
        ctx.uow, ctx.gene_repo, _match_repo(repo, ctx), ctx.dispatcher
    )
    return (await handler(cmd, ctx.auth)).unwrap(), present


async def _run_vulnerability(
    ctx: _KindContext, records: list[Any]
) -> tuple[list[ItemResult], int]:
    repo = SQLAlchemyVulnerabilityRepository(ctx.uow)
    present = await _already_present(repo, ctx)
    cmd = BulkUpsertVulnerabilityCommand(
        target_workspace_id=ctx.target_workspace_id,
        organism_id=ctx.organism_id,
        records=tuple(records),
        dry_run=ctx.dry_run,
    )
    handler = BulkUpsertVulnerability(
        ctx.uow, ctx.gene_repo, _match_repo(repo, ctx), ctx.dispatcher
    )
    return (await handler(cmd, ctx.auth)).unwrap(), present


async def _run_hypomorph(ctx: _KindContext, records: list[Any]) -> tuple[list[ItemResult], int]:
    repo = SQLAlchemyHypomorphRepository(ctx.uow)
    present = await _already_present(repo, ctx)
    # A fresh, unwrapped instance: hypomorph resolves knockdown_strain by name
    # against this repo's find_by_gene — a *foreign* lookup, not hypomorph's own
    # match — which must never be short-circuited by add mode.
    strain_repo = SQLAlchemyCrispriStrainRepository(ctx.uow)
    cmd = BulkUpsertHypomorphCommand(
        target_workspace_id=ctx.target_workspace_id,
        organism_id=ctx.organism_id,
        records=tuple(records),
        dry_run=ctx.dry_run,
    )
    handler = BulkUpsertHypomorph(
        ctx.uow, ctx.gene_repo, _match_repo(repo, ctx), strain_repo, ctx.dispatcher
    )
    return (await handler(cmd, ctx.auth)).unwrap(), present


async def _run_crispri_strain(
    ctx: _KindContext, records: list[Any]
) -> tuple[list[ItemResult], int]:
    repo = SQLAlchemyCrispriStrainRepository(ctx.uow)
    present = await _already_present(repo, ctx)
    cmd = BulkUpsertCrispriStrainCommand(
        target_workspace_id=ctx.target_workspace_id,
        organism_id=ctx.organism_id,
        records=tuple(records),
        dry_run=ctx.dry_run,
    )
    handler = BulkUpsertCrispriStrain(
        ctx.uow, ctx.gene_repo, _match_repo(repo, ctx), ctx.dispatcher
    )
    return (await handler(cmd, ctx.auth)).unwrap(), present


async def _run_resistance_mutation(
    ctx: _KindContext, records: list[Any]
) -> tuple[list[ItemResult], int]:
    repo = SQLAlchemyResistanceMutationRepository(ctx.uow)
    present = await _already_present(repo, ctx)
    cmd = BulkUpsertResistanceMutationCommand(
        target_workspace_id=ctx.target_workspace_id,
        organism_id=ctx.organism_id,
        records=tuple(records),
        dry_run=ctx.dry_run,
    )
    handler = BulkUpsertResistanceMutation(
        ctx.uow, ctx.gene_repo, _match_repo(repo, ctx), ctx.dispatcher
    )
    return (await handler(cmd, ctx.auth)).unwrap(), present


async def _run_protein_production(
    ctx: _KindContext, records: list[Any]
) -> tuple[list[ItemResult], int]:
    repo = SQLAlchemyProteinProductionRepository(ctx.uow)
    present = await _already_present(repo, ctx)
    cmd = BulkUpsertProteinProductionCommand(
        target_workspace_id=ctx.target_workspace_id,
        records=tuple(records),
        dry_run=ctx.dry_run,
    )
    handler = BulkUpsertProteinProduction(
        ctx.uow, ctx.protein_repo, _match_repo(repo, ctx), ctx.dispatcher
    )
    return (await handler(cmd, ctx.auth)).unwrap(), present


async def _run_protein_activity_assay(
    ctx: _KindContext, records: list[Any]
) -> tuple[list[ItemResult], int]:
    repo = SQLAlchemyProteinActivityAssayRepository(ctx.uow)
    present = await _already_present(repo, ctx)
    cmd = BulkUpsertProteinActivityAssayCommand(
        target_workspace_id=ctx.target_workspace_id,
        records=tuple(records),
        dry_run=ctx.dry_run,
    )
    handler = BulkUpsertProteinActivityAssay(
        ctx.uow, ctx.protein_repo, _match_repo(repo, ctx), ctx.dispatcher
    )
    return (await handler(cmd, ctx.auth)).unwrap(), present


async def _run_unpublished_structure(
    ctx: _KindContext, records: list[Any]
) -> tuple[list[ItemResult], int]:
    repo = SQLAlchemyUnpublishedStructureRepository(ctx.uow)
    present = await _already_present(repo, ctx)
    cmd = BulkUpsertUnpublishedStructureCommand(
        target_workspace_id=ctx.target_workspace_id,
        records=tuple(records),
        dry_run=ctx.dry_run,
    )
    handler = BulkUpsertUnpublishedStructure(
        ctx.uow, ctx.protein_repo, _match_repo(repo, ctx), ctx.dispatcher
    )
    return (await handler(cmd, ctx.auth)).unwrap(), present


_DISPATCH: dict[
    RecordKind, Callable[[_KindContext, list[Any]], Awaitable[tuple[list[ItemResult], int]]]
] = {
    RecordKind.ESSENTIALITY: _run_essentiality,
    RecordKind.VULNERABILITY: _run_vulnerability,
    RecordKind.HYPOMORPH: _run_hypomorph,
    RecordKind.CRISPRI_STRAIN: _run_crispri_strain,
    RecordKind.RESISTANCE_MUTATION: _run_resistance_mutation,
    RecordKind.PROTEIN_PRODUCTION: _run_protein_production,
    RecordKind.PROTEIN_ACTIVITY_ASSAY: _run_protein_activity_assay,
    RecordKind.UNPUBLISHED_STRUCTURE: _run_unpublished_structure,
}


_UNMATCHED_PREFIXES = ("unmatched locus ", "unmatched accession ")


def _unmatched_value(error: str) -> str | None:
    """Recovers the offending locus/accession from a bulk command's own failure
    message. Coupled to the exact wording every one of the eight commands uses
    (``f"unmatched locus {x}"`` / ``f"unmatched accession {x}"``) rather than
    re-deriving unmatched-ness independently, so this always agrees with what
    the command itself decided — never a second, possibly-diverging guess.
    """
    for prefix in _UNMATCHED_PREFIXES:
        if error.startswith(prefix):
            return error[len(prefix) :]
    return None


def _row_problem(sheet: str, reason: str) -> dict[str, Any]:
    # row is None: this problem was found after parse-time deduplication, which
    # does not carry the original spreadsheet row number forward onto a record
    # (RowProblem.row exists only for problems parse_workbook itself raises —
    # see SheetPlan vs *ImportRecord in target_biology_workbook.py). Naming the
    # sheet and the exact offending value is the most that can be said here.
    return {"sheet": sheet, "row": None, "reason": reason}


def _drop_ambiguous(
    kind: RecordKind, records: list[Any], index: LocusIndex
) -> tuple[list[Any], list[dict[str, Any]]]:
    """Splits out rows whose locus/name is claimed by more than one gene before
    dispatch. Left in, they would reach the bulk command's own "unmatched"
    check indistinguishably from a genuinely-missing locus — build_locus_index
    already excludes ambiguous keys from what ``.get()`` returns, so the
    command would just say "unmatched locus X" either way. Pulled out here
    instead, the row is failed naming every candidate gene, which the command
    has no way to do (it never sees the ambiguity, only the miss).
    """
    kept: list[Any] = []
    problems: list[dict[str, Any]] = []
    for rec in records:
        candidates = index.ambiguous.get(rec.locus_key.upper())
        if candidates is None:
            kept.append(rec)
            continue
        names = ", ".join(sorted(g.primary_name for g in candidates))
        problems.append(
            _row_problem(
                kind.value,
                f"ambiguous locus {rec.locus_key!r}: matches {len(candidates)} genes ({names})",
            )
        )
    return kept, problems


def _unresolved_ligand_warning(
    kind: RecordKind, records: list[Any], update_existing: bool
) -> dict[str, Any] | None:
    """Update mode's natural key for ``unpublished_structure`` is ``(protein_id,
    method, ligands)``, and ``ligands`` only ever means *resolved* compound ids
    — a cell whose ligand text didn't resolve to one lands in
    ``extensions["ligand_reported"]`` instead (``target_biology_workbook``'s
    module docstring, "Unresolvable ligand text"), invisible to that key. Two
    such rows for the same protein and method are indistinguishable to the
    command and one silently overwrites the other on match — see
    ``bulk_upsert_unpublished_structure.py``'s module docstring, the exact bug
    class this migration exists to close. Add mode never matches at all (see
    ``_NoExistingMatch``), so it can't hit this; only worth a warning when
    update mode is actually selected.
    """
    if kind is not RecordKind.UNPUBLISHED_STRUCTURE or not update_existing:
        return None
    affected = sum(1 for rec in records if (rec.extensions or {}).get("ligand_reported"))
    if affected == 0:
        return None
    return {
        "sheet": kind.value,
        "count": affected,
        "reason": (
            "update mode is selected and this sheet has rows whose ligand text "
            "did not resolve to a compound id — the upsert key cannot "
            "discriminate them by ligand, so a match may silently overwrite a "
            "different structure"
        ),
    }


def _ignored_columns(data: bytes) -> dict[str, list[str]]:
    """Per-sheet headers the parser recognises but drops on the floor — the
    five provenance columns the owner decided are not worth keeping (see
    ``target_biology_workbook``'s module docstring, header group 2: source_type/
    url/note/contributor/observed_on). A second, cheap pass over just each
    sheet's header row — ``parse_workbook`` already read the whole workbook
    once and does not surface which of these it saw, only that it skipped them.
    """
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    kind_names = {k.value for k in RecordKind}
    found: dict[str, list[str]] = {}
    for sheet_name in wb.sheetnames:
        key = sheet_name.strip().casefold()
        if key not in kind_names:
            continue
        header_row = next(wb[sheet_name].iter_rows(values_only=True), ())
        header = {
            str(c).strip().casefold() for c in header_row if c is not None and str(c).strip()
        }
        dropped = sorted(_PROVENANCE_DROPPED & header)
        if dropped:
            found[key] = dropped
    return found


class TargetBiologyAdapter:
    """Wires the eight target-biology bulk-upsert commands from a parsed
    workbook (Task 4's ``parse_workbook``).

    Preview (``dry_run=True``) and apply (``dry_run=False``) are two ``ImportRun``s
    over the *same* stored upload — this adapter does not distinguish them
    beyond passing ``dry_run`` straight through to every bulk command, each of
    which already no-ops its write when it is set; this adapter's own reads
    (``already_present``, the locus index) never write anything, so they are
    preview-safe by construction too.
    """

    import_type = ImportType.TARGET_BIOLOGY

    async def run(self, rt: ImportRuntime) -> dict[str, Any]:
        params = TargetBiologyParams(**rt.params)
        # SECURITY: the target workspace is the run's OWN workspace_id — set at
        # StartImport time from the *caller's* real auth.workspace_id, and
        # threaded here exactly like the reporter's workspace already is. Never
        # params (client-settable) and never rt.auth (ServiceAuth, always
        # SHARED_WORKSPACE_ID — it exists only to satisfy each bulk command's
        # own require_admin check). See ImportRuntime.workspace_id's docstring.
        target_workspace_id = rt.workspace_id

        await rt.reporter.phase("loading upload")
        data = await rt.load_upload(params.upload_ref)

        await rt.reporter.phase("parsing")
        ext_uow = AsyncUnitOfWork(rt.session_factory)
        async with ext_uow:
            field_defs = await SQLAlchemyExtensionFieldDefRepository(ext_uow).list_all(
                target_workspace_id
            )
        known_extension_fields: dict[str, dict[str, str]] = {}
        for field_def in field_defs:
            known_extension_fields.setdefault(field_def.kind, {})[field_def.name] = (
                field_def.field_type.value
            )

        plans, workbook_problems = parse_workbook(
            data, match_by=params.match_by, known_extension_fields=known_extension_fields
        )

        problems: list[dict[str, Any]] = [
            _row_problem(p.sheet, p.reason) for p in workbook_problems
        ]
        unmatched: set[str] = set()
        kinds_summary: dict[str, dict[str, int]] = {}
        already_present: dict[str, int] = {}
        warnings: list[dict[str, Any]] = []

        uow = AsyncUnitOfWork(rt.session_factory)
        ctx = _KindContext(
            uow=uow,
            dispatcher=rt.dispatcher,
            auth=rt.auth,
            target_workspace_id=target_workspace_id,
            organism_id=params.organism_id,
            dry_run=params.dry_run,
            update_existing=params.update_existing,
            gene_repo=SQLAlchemyGeneRepository(uow),
            protein_repo=SQLAlchemyProteinRepository(uow),
        )

        # One shared locus index for every gene-side sheet in this workbook —
        # each bulk command below still rebuilds its own internally (it has no
        # way to accept one pre-built), but this pass exists only to catch
        # ambiguity before dispatch, which the commands cannot report at all.
        locus_index: LocusIndex | None = None
        if any(plan.kind in _GENE_SIDE_KINDS for plan in plans):
            async with AsyncUnitOfWork(rt.session_factory) as genes_uow:
                genes = await SQLAlchemyGeneRepository(genes_uow).list_by_organism(
                    params.organism_id, workspace_id=target_workspace_id
                )
            locus_index = build_locus_index(genes)

        total_rows = sum(plan.rows_read for plan in plans)
        processed = 0
        await rt.reporter.advance(processed, total_rows)

        for plan in plans:
            await rt.reporter.phase(f"importing {plan.kind.value}")
            records = plan.records
            plan_problems = [_row_problem(plan.kind.value, p.reason) for p in plan.problems]
            failed = len(plan.problems)

            if plan.kind in _GENE_SIDE_KINDS and locus_index is not None:
                records, ambiguous_problems = _drop_ambiguous(plan.kind, records, locus_index)
                failed += len(ambiguous_problems)
                plan_problems.extend(ambiguous_problems)

            ligand_warning = _unresolved_ligand_warning(plan.kind, records, ctx.update_existing)
            if ligand_warning is not None:
                warnings.append(ligand_warning)

            results, present = await _DISPATCH[plan.kind](ctx, records)
            already_present[plan.kind.value] = present

            created = updated = 0
            for result in results:
                if result.status == "created":
                    created += 1
                elif result.status == "updated":
                    updated += 1
                elif result.status == "failed":
                    failed += 1
                    reason = result.error or "failed"
                    value = _unmatched_value(reason)
                    if value is not None:
                        unmatched.add(value)
                    plan_problems.append(_row_problem(plan.kind.value, reason))

            kinds_summary[plan.kind.value] = {
                "rows": plan.rows_read,
                "records": len(plan.records),
                "merged_identical": plan.merged_identical,
                "create": created,
                "update": updated,
                "failed": failed,
            }
            problems.extend(plan_problems)
            processed += plan.rows_read
            await rt.reporter.advance(processed, total_rows)

        return {
            "kinds": kinds_summary,
            "unmatched": {"count": len(unmatched), "examples": sorted(unmatched)[:50]},
            "problems": problems[:50],
            "problems_truncated": max(0, len(problems) - 50),
            "already_present": already_present,
            "ignored_columns": _ignored_columns(data),
            "warnings": warnings,
        }


# ---------------------------------------------------------------------------
# Registry — the worker looks up adapters here
# ---------------------------------------------------------------------------

IMPORT_ADAPTERS: dict[ImportType, ImportAdapter] = {
    ImportType.PROTEOME: ProteomeAdapter(),
    ImportType.GENE_ENRICHMENT: GeneEnrichmentAdapter(),
    ImportType.GO_ONTOLOGY: GoOntologyAdapter(),
    ImportType.PLUGIN: PluginDispatchAdapter(),
    ImportType.TARGET_BIOLOGY: TargetBiologyAdapter(),
}
